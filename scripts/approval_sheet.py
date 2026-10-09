# -*- coding: utf-8 -*-
"""Phiếu duyệt từng dòng bằng Excel — dùng khi không có AI hoặc AI không có hộp hỏi đáp.

review_spec.py ghi phiếu (write_sheet); người dùng chọn cột "Quyết định" cho TỪNG dòng; apply_fixes.py đọc lại (read_sheet).
Quyết định:
  - "Sửa theo đề xuất" : dùng giá trị ở cột Đề xuất
  - "Giá trị khác"     : dùng giá trị người dùng gõ ở cột "Giá trị khác"
  - "Giữ nguyên" / để trống : KHÔNG đổi (mặc định an toàn)
"""
FIX, KEEP, OTHER = 'Sửa theo đề xuất', 'Giữ nguyên', 'Giá trị khác'
HEAD = ['No.', 'Mã', 'Form', 'Vị trí', 'Hiện tại', 'Đề xuất', 'Lý do', 'Quyết định', 'Giá trị khác']
GUIDE = [
    'CÁCH DUYỆT',
    '1. Sang sheet "Duyet". Mỗi dòng là 1 phát hiện cần bạn quyết định.',
    f'2. Ở cột "Quyết định", chọn: "{FIX}", "{KEEP}" hoặc "{OTHER}".',
    f'3. Chọn "{OTHER}" thì gõ giá trị đúng vào cột "Giá trị khác".',
    f'4. Dòng để trống Quyết định = "{KEEP}". Không có dòng nào bị đổi nếu bạn không chọn.',
    '5. Không sửa các cột Mã, Hiện tại, Đề xuất. Lưu file rồi chạy apply_fixes.py với file này.',
    'Dòng không có Đề xuất là câu hỏi cần xác nhận: chỉ chọn được Giữ nguyên hoặc Giá trị khác.',
]


def write_sheet(findings, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = 'Duyet'
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='D9E1F2')
    dv_fix = DataValidation(type='list', formula1=f'"{FIX},{KEEP},{OTHER}"', allow_blank=True)
    dv_ask = DataValidation(type='list', formula1=f'"{KEEP},{OTHER}"', allow_blank=True)
    ws.add_data_validation(dv_fix)
    ws.add_data_validation(dv_ask)
    for i, x in enumerate(findings, 1):
        ws.append([i, x['id'], x['form'], x['where'], x['current'], x.get('suggestion') or '(cần xác nhận)',
                   '; '.join(x.get('reasons') or []), '', ''])
        (dv_fix if x.get('suggestion') else dv_ask).add(f'H{i + 1}')
        for col in 'HI':
            ws[f'{col}{i + 1}'].fill = PatternFill('solid', fgColor='FFF2CC')
    for col, w in zip('ABCDEFGHI', (5, 7, 8, 38, 38, 38, 50, 18, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'C2'
    g = wb.create_sheet('Huong dan')
    for line in GUIDE:
        g.append([line])
    g['A1'].font = Font(bold=True)
    g.column_dimensions['A'].width = 110
    wb.save(path)


def read_sheet(path):
    """Trả về (approvals, warnings). approvals cùng dạng approvals.json: {id: "approve" | "reject" | {"value": ...}}."""
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)['Duyet']
    head = [str(c.value or '').strip() for c in ws[1]]
    ix = {h: head.index(h) for h in HEAD if h in head}
    missing = [h for h in ('Mã', 'Quyết định', 'Giá trị khác') if h not in ix]
    if missing:
        raise SystemExit(f'{path}: thiếu cột {", ".join(missing)} trong sheet "Duyet"')
    appr, warn = {}, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        fid = str(row[ix['Mã']] or '').strip()
        if not fid:
            continue
        dec = str(row[ix['Quyết định']] or '').strip()
        val = row[ix['Giá trị khác']]
        val = '' if val is None else str(val)
        if dec == FIX:
            appr[fid] = 'approve'
        elif dec == OTHER:
            if val.strip():
                appr[fid] = {'value': val.strip()}
            else:
                appr[fid] = 'reject'
                warn.append(f'{fid}: chọn "{OTHER}" nhưng chưa gõ giá trị → giữ nguyên')
        elif dec in ('', KEEP):
            appr[fid] = 'reject'
            if val.strip():
                warn.append(f'{fid}: có "Giá trị khác" nhưng chưa chọn Quyết định → giữ nguyên')
        else:
            appr[fid] = 'reject'
            warn.append(f'{fid}: Quyết định "{dec}" không hợp lệ → giữ nguyên')
    return appr, warn
