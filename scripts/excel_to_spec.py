# -*- coding: utf-8 -*-
"""Đọc 1 sheet Excel dạng bảng phẳng -> spec Eform (cột phẳng). Cây cột nhóm dựng thêm sau (xem SKILL.md).

Dùng:
  python3 excel_to_spec.py file.xlsx --sheet Tinh --key tinh --group "Biểu mẫu Kiểm soát TTHC" \
      --attr TTHC_CapTinh --fixed "B,C,D" [--name-cell B1] [--header-row 2] [--data-start 3] \
      [--skip "A"] [--spec spec.json]
- Tên E-Form lấy từ ô --name-cell (bỏ tiền tố "Tên eform:").
- Cột trong --fixed: KIỂU Text, giá trị từng dòng thành dữ liệu cố định.
- Cột còn lại (trừ --skip, mặc định cột A = STT vì bảng tự đánh STT): KIỂU Number, để người nhập điền.
- Nếu --spec đã tồn tại, form mới được thêm/ghi đè vào khóa --key.
"""
import argparse, json, os, sys
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import utf8_console
from openpyxl.utils import column_index_from_string, get_column_letter


def main():
    utf8_console()
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx'); ap.add_argument('--sheet', required=True); ap.add_argument('--key', required=True)
    ap.add_argument('--group', default=''); ap.add_argument('--attr', required=True)
    ap.add_argument('--title', default=''); ap.add_argument('--heading', default='')
    ap.add_argument('--fixed', default='B'); ap.add_argument('--skip', default='A')
    ap.add_argument('--name-cell', default='B1'); ap.add_argument('--header-row', type=int, default=2)
    ap.add_argument('--data-start', type=int, default=3); ap.add_argument('--spec', default='spec.json')
    ap.add_argument('--text-width', type=int, default=300); ap.add_argument('--num-width', type=int, default=150)
    a = ap.parse_args()
    ws = openpyxl.load_workbook(a.xlsx, data_only=True)[a.sheet]
    name = str(ws[a.name_cell].value or '').strip()
    if ':' in name and name.lower().startswith('tên eform'):
        name = name.split(':', 1)[1].strip()
    fixed = {c.strip().upper() for c in a.fixed.split(',') if c.strip()}
    skip = {c.strip().upper() for c in a.skip.split(',') if c.strip()}
    rows = [r for r in range(a.data_start, ws.max_row + 1)
            if any(ws.cell(r, column_index_from_string(c)).value not in (None, '') for c in fixed)]
    cols, seen = [], {}
    for ci in range(1, ws.max_column + 1):
        L = get_column_letter(ci)
        h = ws.cell(a.header_row, ci).value
        if L in skip or h in (None, ''):
            continue
        h = ' '.join(str(h).split())
        seen[h] = seen.get(h, 0) + 1
        if L in fixed:
            vals = [' '.join(str(ws.cell(r, ci).value if ws.cell(r, ci).value is not None else '').split()) for r in rows]
            cols.append({'title': h, 'attr': h, 'type': 'Text', 'width': a.text_width if L == min(fixed) else 80, 'fixed': vals, 'excel_col': L})
        else:
            cols.append({'title': h, 'attr': h, 'type': 'Number', 'width': a.num_width, 'excel_col': L})
    dups = [h for h, n in seen.items() if n > 1]
    spec = json.load(open(a.spec, encoding='utf-8')) if os.path.exists(a.spec) else {'source': os.path.basename(a.xlsx), 'forms': {}}
    spec['forms'][a.key] = {'sheet': a.sheet, 'name': name, 'group': a.group, 'attr': a.attr, 'title': a.title or name,
                            'heading': a.heading or name.upper(), 'fixed_row': True, 'sticky_header': True, 'columns': cols}
    json.dump(spec, open(a.spec, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'{a.key}: "{name}" | {len(rows)} dòng | {len(cols)} cột -> {a.spec}')
    if dups:
        print('CẢNH BÁO tiêu đề cột trùng (cần hỏi người dùng / đối chiếu biểu chuẩn):', dups)


if __name__ == '__main__':
    main()
